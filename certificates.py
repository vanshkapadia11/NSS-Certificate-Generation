"""Excel ingestion, PDF rendering and durable SMTP delivery tracking."""
import hashlib
import io
import re
import smtplib
import sqlite3
import ssl
from functools import lru_cache
from email.message import EmailMessage
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import landscape, A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase.pdfmetrics import stringWidth


def read_participants(source, sheet=None):
    book = load_workbook(source, read_only=True, data_only=True)
    try:
        ws = book[sheet] if sheet else book.worksheets[0]
        rows = iter(ws.iter_rows(values_only=True))
        headers = [str(v or '').strip().lower() for v in next(rows, [])]
        aliases = {'name': ['name', 'participant name', 'full name'],
                   'email': ['email', 'email address', 'email id']}
        columns = {}
        for field, choices in aliases.items():
            matches = [i for i, h in enumerate(headers) if h in choices]
            if len(matches) != 1:
                raise ValueError(f'Provide exactly one {field} column in the first row.')
            columns[field] = matches[0]
        valid, errors, seen = [], [], set()
        for number, row in enumerate(rows, 2):
            if not any(v is not None for v in row):
                continue
            name = str(row[columns['name']] or '').strip()
            email = str(row[columns['email']] or '').strip()
            if not name or len(name) > 120 or any(ord(c) < 32 for c in name):
                errors.append({'row': number, 'problem': 'Name is missing, too long, or contains control characters.'})
                continue
            if not re.fullmatch(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9](?:[A-Za-z0-9.-]*[A-Za-z0-9])?\.[A-Za-z]{2,}", email):
                errors.append({'row': number, 'problem': 'Invalid email address.'})
                continue
            key = email.lower()
            if key in seen:
                errors.append({'row': number, 'problem': 'Duplicate email address; keep one participant per email.'})
                continue
            seen.add(key)
            valid.append({'name': name, 'email': email})
        return valid, errors
    finally:
        book.close()


def certificate_id(event_id, email):
    return hashlib.sha256(f'{event_id.strip()}|{email.lower()}'.encode()).hexdigest()[:16].upper()


def make_pdf(person, config):
    buffer = io.BytesIO()
    w, h = landscape(A4)
    c = canvas.Canvas(buffer, pagesize=(w, h))
    c.setTitle(f"Grainathon certificate - {person['name']}")
    green, gold = HexColor('#163F34'), HexColor('#B58C42')
    c.setFillColor(HexColor('#FBF8EF')); c.rect(0, 0, w, h, fill=1, stroke=0)
    c.setStrokeColor(green); c.setLineWidth(3); c.rect(22, 22, w-44, h-44)
    c.setStrokeColor(gold); c.setLineWidth(0.7); c.rect(31, 31, w-62, h-62)
    # Stylized ears of grain drawn as vectors, crisp at print resolution.
    for x, direction in [(72, 1), (w-72, -1)]:
        c.setStrokeColor(gold); c.setFillColor(gold); c.setLineWidth(2)
        c.line(x, 150, x + direction*17, 410)
        for y in range(205, 400, 22):
            stem = x + direction*(y-150)*17/260
            for side in [-1, 1]:
                c.saveState(); c.translate(stem, y); c.rotate(side*42)
                c.ellipse(-4, 0, 4, 22, fill=1, stroke=0); c.restoreState()

    def centered(text, y, size, font='Helvetica', color=green, width=620):
        text = str(text)
        while stringWidth(text, font, size) > width and size > 8:
            size -= 0.5
        c.setFillColor(color); c.setFont(font, size); c.drawCentredString(w/2, y, text)

    centered(config['organizer'], 520, 15, 'Helvetica-Bold')
    centered('G R A I N A T H O N', 460, 34, 'Times-Bold')
    centered('CERTIFICATE OF PARTICIPATION', 422, 15, color=gold)
    centered('This certificate is proudly presented to', 367, 13)
    centered(person['name'], 310, 36, 'Times-Bold', width=590)
    c.setStrokeColor(gold); c.line(180, 291, w-180, 291)
    centered('for participating in Grainathon', 260, 17)
    centered(config['tagline'], 233, 12)
    centered(f"Held on {config['date']}  |  {config['venue']}", 207, 12)
    for x, name, role in [(245, config['signatory1'], config['role1']),
                           (w-245, config['signatory2'], config['role2'])]:
        c.setStrokeColor(gold); c.line(x-92, 128, x+92, 128)
        for text, y, font in [(name, 110, 'Helvetica-Bold'), (role, 93, 'Helvetica')]:
            size = 11
            while stringWidth(text, font, size) > 205 and size > 7: size -= .5
            c.setFillColor(green); c.setFont(font, size); c.drawCentredString(x, y, text)
    centered('Certificate ID: ' + certificate_id(config['event_id'], person['email']), 50, 8, color=gold)
    c.showPage(); c.save()
    return buffer.getvalue()


@lru_cache(maxsize=1)
def prepared_template(template, filename):
    import fitz
    source = fitz.open(stream=template, filetype=filename.rsplit('.', 1)[-1].lower())
    if not source.is_pdf:
        converted = source.convert_to_pdf()
        source.close()
        source = fitz.open(stream=converted, filetype='pdf')
    if len(source) != 1:
        source.close()
        raise ValueError('Please use a single-page certificate template.')
    font_path = Path(__file__).parent / 'assets/fonts/Poppins-Medium.ttf'
    font = fitz.Font(fontfile=str(font_path))
    source[0].insert_font(fontname='ParticipantFont', fontfile=str(font_path))
    data = source.tobytes()
    source.close()
    return data, font


def overlay_name(template, filename, name, x_percent, y_percent, size, color):
    import fitz
    data, font = prepared_template(template, filename)
    source = fitz.open(stream=data, filetype='pdf')
    page = source[0]
    width, height = page.rect.width, page.rect.height
    fontname = 'ParticipantFont'
    if any(font.has_glyph(ord(ch)) == 0 for ch in name):
        raise ValueError('The name contains characters unsupported by the installed font.')
    x, y = width*x_percent/100, height*y_percent/100
    available = 2*min(x-15, width-x-15)
    text_width = font.text_length(name, fontsize=size)
    if text_width > available:
        size *= available/text_width
        text_width = font.text_length(name, fontsize=size)
    if size < 8 or y-size < 0 or y > height-10:
        raise ValueError('Name position leaves insufficient space. Move it towards the center.')
    rgb = tuple(int(color[i:i+2],16)/255 for i in (1,3,5))
    page.insert_text((x-text_width/2, y), name, fontsize=size, fontname=fontname, color=rgb)
    result = source.tobytes(deflate=True)
    source.close()
    return result


def connect_log(path='data/delivery.sqlite3'):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(path)
    db.execute('CREATE TABLE IF NOT EXISTS deliveries (id TEXT PRIMARY KEY, email TEXT, status TEXT, updated TEXT DEFAULT CURRENT_TIMESTAMP)')
    db.execute('CREATE TABLE IF NOT EXISTS delivery_history (id TEXT, email TEXT, status TEXT, updated TEXT)')
    db.commit()
    return db


class SMTPBatch:
    """Lazily authenticate once, then reuse the connection for the batch."""
    def __init__(self, settings):
        self.settings = settings
        self.client = None

    def send_message(self, message):
        if self.client is None:
            smtp = self.settings
            context = ssl.create_default_context()
            client = smtplib.SMTP_SSL(smtp['host'], smtp['port'], timeout=30, context=context) if smtp['ssl'] else smtplib.SMTP(smtp['host'], smtp['port'], timeout=30)
            try:
                if not smtp['ssl']:
                    client.ehlo(); client.starttls(context=context); client.ehlo()
                client.login(smtp['username'], smtp['password'])
            except Exception:
                client.close()
                raise
            self.client = client
        self.client.send_message(message)

    def close(self):
        if self.client is not None:
            try:
                self.client.quit()
            except Exception:
                self.client.close()
            self.client = None


def send_certificate(person, config, pdf, smtp, db, transport=None, resend=False):
    key = certificate_id(config['event_id'], person['email'])
    message = EmailMessage()
    message['From'] = smtp['sender']; message['To'] = person['email']
    message['Subject'] = 'Your Grainathon participation certificate'
    message.set_content(f"Dear {person['name']},\n\nThank you for participating in Grainathon! Your certificate is attached.\n\nRegards,\n{config['organizer']}")
    message.add_attachment(pdf, maintype='application', subtype='pdf', filename=f'Grainathon-{key}.pdf')
    # Atomic reservation prevents concurrent script runs from double-sending.
    with db:
        cursor = db.execute('INSERT OR IGNORE INTO deliveries(id,email,status) VALUES(?,?,?)', (key, person['email'], 'pending'))
        if cursor.rowcount == 0 and resend:
            db.execute('INSERT INTO delivery_history SELECT id,email,status,updated FROM deliveries WHERE id=? AND status=?', (key, 'accepted by SMTP'))
            cursor = db.execute('UPDATE deliveries SET status=?, updated=CURRENT_TIMESTAMP WHERE id=? AND status=?', ('pending', key, 'accepted by SMTP'))
    if cursor.rowcount == 0:
        return 'skipped (already recorded; use --send --resend for previously accepted deliveries)'
    connection = transport if transport is not None else SMTPBatch(smtp)
    try:
        connection.send_message(message)
        status = 'accepted by SMTP'
    except Exception:
        # A timeout may happen after acceptance: never blindly retry.
        status = 'uncertain — check provider before retrying'
    finally:
        if transport is None:
            connection.close()
    with db:
        db.execute('UPDATE deliveries SET status=?, updated=CURRENT_TIMESTAMP WHERE id=?', (status, key))
    return status
